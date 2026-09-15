"""Extract text from uploaded documents and prepare image payloads."""

from __future__ import annotations

import base64
import io
from dataclasses import dataclass
from pypdf import PdfReader

# Soft cap for document text sent to the model (chars). Larger docs are truncated.
MAX_DOC_CHARS = 80_000

TEXT_EXTENSIONS = {
    ".txt",
    ".md",
    ".markdown",
    ".csv",
    ".json",
    ".py",
    ".js",
    ".ts",
    ".tsx",
    ".jsx",
    ".html",
    ".htm",
    ".css",
    ".xml",
    ".yaml",
    ".yml",
    ".toml",
    ".ini",
    ".cfg",
    ".log",
    ".rst",
    ".sh",
    ".bash",
    ".sql",
    ".r",
    ".go",
    ".rs",
    ".java",
    ".c",
    ".cpp",
    ".h",
    ".hpp",
}

IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg"}
IMAGE_MIME = {
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
}

VIDEO_EXTENSIONS = {
    ".mp4",
    ".mov",
    ".avi",
    ".mkv",
    ".webm",
    ".m4v",
    ".wmv",
    ".flv",
}


@dataclass
class ExtractedDocument:
    """Text extracted from an uploaded document."""

    name: str
    text: str
    truncated: bool = False
    error: str | None = None


@dataclass
class ExtractedImage:
    """Image payload for vision-capable chat completions."""

    name: str
    mime: str
    data_url: str  # data:<mime>;base64,...
    error: str | None = None


def _extension(name: str) -> str:
    if "." not in name:
        return ""
    return "." + name.rsplit(".", 1)[-1].lower()


def _truncate(text: str, limit: int = MAX_DOC_CHARS) -> tuple[str, bool]:
    if len(text) <= limit:
        return text, False
    note = (
        f"\n\n[Note: document truncated from {len(text):,} to {limit:,} characters "
        "to stay within context limits.]"
    )
    keep = max(0, limit - len(note))
    return text[:keep] + note, True


def _read_pdf(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    parts: list[str] = []
    for i, page in enumerate(reader.pages, start=1):
        page_text = page.extract_text() or ""
        if page_text.strip():
            parts.append(f"--- Page {i} ---\n{page_text}")
    return "\n\n".join(parts).strip()


def _read_docx(data: bytes) -> str:
    from docx import Document  # lazy import

    doc = Document(io.BytesIO(data))
    return "\n".join(p.text for p in doc.paragraphs if p.text).strip()


def _decode_utf8(data: bytes) -> str:
    return data.decode("utf-8")


def extract_document(name: str, data: bytes) -> ExtractedDocument:
    """Extract text from a document upload. Returns error field on failure."""
    ext = _extension(name)

    if ext in VIDEO_EXTENSIONS:
        return ExtractedDocument(
            name=name,
            text="",
            error=(
                f"Uploaded video files are out of scope for this app. "
                f"'{name}' was skipped. Paste a YouTube link instead, or use documents/images."
            ),
        )

    try:
        if ext == ".pdf":
            text = _read_pdf(data)
            if not text:
                return ExtractedDocument(
                    name=name,
                    text="",
                    error=f"Could not extract text from PDF '{name}' (empty or image-only).",
                )
        elif ext == ".docx":
            text = _read_docx(data)
            if not text:
                return ExtractedDocument(
                    name=name,
                    text="",
                    error=f"DOCX '{name}' had no extractable paragraph text.",
                )
        elif ext in TEXT_EXTENSIONS:
            text = _decode_utf8(data)
        else:
            # Unknown extension: try UTF-8 text; fail clearly if binary.
            try:
                text = _decode_utf8(data)
            except UnicodeDecodeError:
                return ExtractedDocument(
                    name=name,
                    text="",
                    error=(
                        f"Could not parse '{name}' "
                        f"(unsupported or binary type '{ext or 'unknown'}')."
                    ),
                )
            if "\x00" in text:
                return ExtractedDocument(
                    name=name,
                    text="",
                    error=f"Could not parse '{name}' (looks like a binary file).",
                )
    except Exception as exc:  # noqa: BLE001 — surface friendly parse errors
        return ExtractedDocument(
            name=name,
            text="",
            error=f"Failed to read '{name}': {exc}",
        )

    text, truncated = _truncate(text)
    return ExtractedDocument(name=name, text=text, truncated=truncated)


def extract_image(name: str, data: bytes, mime_hint: str | None = None) -> ExtractedImage:
    """Build a base64 data-URL image payload for the OpenAI-compatible vision API."""
    ext = _extension(name)
    if ext in VIDEO_EXTENSIONS:
        return ExtractedImage(
            name=name,
            mime="",
            data_url="",
            error=f"Uploaded video '{name}' is out of scope. Use images (PNG/JPEG) or a YouTube link.",
        )

    mime = IMAGE_MIME.get(ext)
    if not mime and mime_hint and mime_hint.startswith("image/"):
        # Only accept officially supported types for xAI image understanding.
        if mime_hint in ("image/png", "image/jpeg", "image/jpg"):
            mime = "image/jpeg" if mime_hint == "image/jpg" else mime_hint

    if mime not in ("image/png", "image/jpeg"):
        return ExtractedImage(
            name=name,
            mime="",
            data_url="",
            error=(
                f"Unsupported image type for '{name}'. "
                "xAI image understanding accepts PNG or JPEG."
            ),
        )

    if len(data) > 20 * 1024 * 1024:
        return ExtractedImage(
            name=name,
            mime=mime,
            data_url="",
            error=f"Image '{name}' exceeds the 20 MiB limit.",
        )

    b64 = base64.standard_b64encode(data).decode("ascii")
    return ExtractedImage(
        name=name,
        mime=mime,
        data_url=f"data:{mime};base64,{b64}",
    )


def is_image_upload(name: str, mime: str | None = None) -> bool:
    ext = _extension(name)
    if ext in IMAGE_EXTENSIONS:
        return True
    if mime and mime.startswith("image/"):
        return True
    return False


def format_documents_for_prompt(docs: list[ExtractedDocument]) -> str:
    """Render extracted documents as a context block for the model."""
    blocks: list[str] = []
    for doc in docs:
        if doc.error or not doc.text:
            continue
        blocks.append(f"### File: {doc.name}\n\n{doc.text}")
    if not blocks:
        return ""
    return (
        "The user attached the following document(s). Use them to answer the instruction.\n\n"
        + "\n\n".join(blocks)
    )
